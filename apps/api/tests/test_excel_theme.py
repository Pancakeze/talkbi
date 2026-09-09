import io
import os

import pandas as pd
import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient
from openpyxl import Workbook

from app.api.routes import data_sources as data_sources_routes
from app.db.session import SessionLocal, engine
from app.models import DataSource, ThemeField, ThemeLibrary, User
from app.services.chat_service import run_chat_query
from app.services.excel_service import (
    DEFAULT_TO_SQL_CHUNKSIZE,
    MAX_SQL_IDENT_LEN,
    PG_MAX_BIND_PARAMS,
    _sanitize_dataframe_columns,
    _to_sql_chunksize,
    materialize_excel_staging,
)


def _xlsx_bytes() -> bytes:
    buf = io.BytesIO()
    df = pd.DataFrame(
        {
            "district_name": ["A区", "B区"],
            "congestion_index": [1.2, 3.4],
            "accident_count": [10, 20],
        }
    )
    with pd.ExcelWriter(buf, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="Traffic", index=False)
    return buf.getvalue()


def test_excel_upload_creates_staging_tables(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("traffic.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    r = client.post("/api/data-sources/excel/upload", files=files, headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["source_type"] == "excel"
    assert body["status"] == "active"
    info = body["connection_info"]
    assert "staging" in info
    assert len(info["staging"]["tables"]) >= 1
    table = info["staging"]["tables"][0]["table"]
    assert table.startswith(f"ds_{body['id']}_")

    detail = client.get(f"/api/data-sources/{body['id']}", headers=headers)
    assert detail.status_code == 200
    assert detail.json()["connection_info"]["staging"]["tables"][0]["row_count"] == 2


def _xlsx_with_empty_extra_sheets() -> bytes:
    """Typical Excel workbook: data on Sheet1 plus unused empty Sheet2/Sheet3."""
    buf = io.BytesIO()
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "district"
    ws1["B1"] = "amount"
    ws1["A2"] = "A区"
    ws1["B2"] = 10.5
    wb.create_sheet("Sheet2")
    wb.create_sheet("Sheet3")
    wb.save(buf)
    return buf.getvalue()


def test_materialize_skips_empty_excel_sheets():
    """
    Empty extra sheets make pandas return a 0-column frame. to_sql then emits
    `CREATE TABLE t ()`, which is invalid SQL and used to roll back the whole
    upload — including sheets that already had data.
    """
    from sqlalchemy import text

    info = materialize_excel_staging(engine, 701, _xlsx_with_empty_extra_sheets(), "classic.xlsx")
    tables = info["staging"]["tables"]
    assert [t["sheet_name"] for t in tables] == ["Sheet1"]
    assert tables[0]["row_count"] == 1
    assert [c["name"] for c in tables[0]["columns"]] == ["district", "amount"]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT district, amount FROM {tables[0]['qualified']}")
            ).mappings()
        ]
    assert rows == [{"district": "A区", "amount": 10.5}]


def test_excel_upload_skips_empty_extra_sheets(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {
        "file": (
            "classic.xlsx",
            _xlsx_with_empty_extra_sheets(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post("/api/data-sources/excel/upload", files=files, headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "active"
    tables = body["connection_info"]["staging"]["tables"]
    assert [t["sheet_name"] for t in tables] == ["Sheet1"]
    assert tables[0]["row_count"] == 1


def test_excel_upload_rejects_workbook_with_only_empty_sheets(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    buf = io.BytesIO()
    Workbook().save(buf)
    files = {
        "file": (
            "empty.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post(
        "/api/data-sources/excel/upload",
        files=files,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 400, r.text
    assert r.json()["detail"] == "No data found in workbook."
    listed = client.get("/api/data-sources", headers={"Authorization": f"Bearer {token}"})
    empty_rows = [ds for ds in listed.json() if ds["name"] == "empty.xlsx"]
    assert empty_rows
    assert empty_rows[-1]["status"] == "failed"


def test_excel_upload_rejects_empty_csv(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    files = {"file": ("empty.csv", b"", "text/csv")}
    r = client.post(
        "/api/data-sources/excel/upload",
        files=files,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 400, r.text
    assert r.json()["detail"] == "No data found in workbook."


def _xlsx_with_far_stray_column(*, extra_empty_sheet: bool = False) -> bytes:
    """Workbook whose used range stretches to column 2001 (past SQLite's 2000 cap)."""
    buf = io.BytesIO()
    wb = Workbook()
    ws1 = wb.active
    ws1.title = "Sheet1"
    ws1["A1"] = "district"
    ws1["B1"] = "amount"
    ws1["A2"] = "A区"
    ws1["B2"] = 10.5
    ws1.cell(1, 2001, "oops")
    if extra_empty_sheet:
        ws2 = wb.create_sheet("Sheet2")
        ws2.cell(1, 2001, "stray")
    wb.save(buf)
    return buf.getvalue()


def test_materialize_drops_empty_unnamed_columns_from_used_range_bloat():
    """
    A stray cell in column 2001 expands Excel's used range. pandas then emits
    ~2000 empty Unnamed columns. to_sql CREATE TABLE exceeds SQLite's default
    2000-column limit (PostgreSQL: 1600) and used to roll back the upload —
    including the two real columns that had data.
    """
    from sqlalchemy import text

    info = materialize_excel_staging(
        engine, 801, _xlsx_with_far_stray_column(), "bloated.xlsx"
    )
    tables = info["staging"]["tables"]
    assert [t["sheet_name"] for t in tables] == ["Sheet1"]
    names = [c["name"] for c in tables[0]["columns"]]
    assert names == ["district", "amount", "oops"]
    assert tables[0]["row_count"] == 1
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT district, amount, oops FROM {tables[0]['qualified']}")
            ).mappings()
        ]
    assert rows == [{"district": "A区", "amount": 10.5, "oops": None}]


def test_materialize_far_column_on_extra_sheet_does_not_roll_back_data():
    """Sheet2 used-range bloat must not abort CREATE TABLE for Sheet1."""
    from sqlalchemy import text

    info = materialize_excel_staging(
        engine,
        802,
        _xlsx_with_far_stray_column(extra_empty_sheet=True),
        "mixed-bloat.xlsx",
    )
    by_sheet = {t["sheet_name"]: t for t in info["staging"]["tables"]}
    assert "Sheet1" in by_sheet
    assert [c["name"] for c in by_sheet["Sheet1"]["columns"]] == [
        "district",
        "amount",
        "oops",
    ]
    with engine.connect() as conn:
        rows = list(
            conn.execute(
                text(f"SELECT district, amount FROM {by_sheet['Sheet1']['qualified']}")
            ).mappings()
        )
    assert [dict(r) for r in rows] == [{"district": "A区", "amount": 10.5}]


def test_materialize_keeps_unnamed_columns_that_have_values():
    """A missing header over real data must not be dropped as used-range bloat."""
    from sqlalchemy import text

    buf = io.BytesIO()
    wb = Workbook()
    ws = wb.active
    ws["A1"] = None
    ws["B1"] = "amount"
    ws["A2"] = "A区"
    ws["B2"] = 10.5
    wb.save(buf)
    info = materialize_excel_staging(engine, 803, buf.getvalue(), "blank-header.xlsx")
    names = [c["name"] for c in info["staging"]["tables"][0]["columns"]]
    assert len(names) == 2
    assert "amount" in names
    unnamed = [n for n in names if n != "amount"][0]
    with engine.connect() as conn:
        row = (
            conn.execute(
                text(
                    f'SELECT "{unnamed}", amount FROM {info["staging"]["tables"][0]["qualified"]}'
                )
            )
            .mappings()
            .one()
        )
    assert dict(row)[unnamed] == "A区"
    assert dict(row)["amount"] == 10.5


def test_excel_upload_survives_used_range_bloat(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    files = {
        "file": (
            "bloated.xlsx",
            _xlsx_with_far_stray_column(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post(
        "/api/data-sources/excel/upload",
        files=files,
        headers={"Authorization": f"Bearer {token}"},
    )
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "active"
    names = [c["name"] for c in body["connection_info"]["staging"]["tables"][0]["columns"]]
    assert names == ["district", "amount", "oops"]


def test_theme_linked_to_datasource_and_chat_uses_staging(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    files = {"file": ("t.xlsx", _xlsx_bytes(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    up = client.post("/api/data-sources/excel/upload", files=files, headers=headers)
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "语义测试库-1", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]

    for col in ("district_name", "congestion_index", "accident_count"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={
                "table_name": physical,
                "field_name": col,
                "alias_zh": col,
                "visible": True,
            },
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    patch = client.patch(
        f"/api/theme-libraries/{theme_id}/fields/{fr.json()['id']}",
        json={"alias_zh": "事故数"},
        headers=headers,
    )
    assert patch.status_code == 200
    assert patch.json()["alias_zh"] == "事故数"

    jr = client.post(
        f"/api/theme-libraries/{theme_id}/joins",
        json={
            "left_table": physical,
            "right_table": physical,
            "left_column": "district_name",
            "right_column": "district_name",
            "join_type": "inner",
        },
        headers=headers,
    )
    assert jr.status_code == 200
    joins = client.get(f"/api/theme-libraries/{theme_id}/joins", headers=headers)
    assert len(joins.json()) == 1

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看各区指标", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    data = chat.json()
    assert data["status"] == "success"
    assert "SELECT" in data["sql"].upper()
    assert len(data["rows"]) == 2
    assert "district_name" in data["rows"][0]


def test_chat_ignores_client_forged_staging_metadata(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    forged = client.post(
        "/api/data-sources",
        json={
            "name": "forged-staging",
            "source_type": "excel",
            "connection_info": {
                "staging": {
                    "dialect": "sqlite",
                    "schema": None,
                    "tables": [
                        {
                            "table": "users",
                            "schema": None,
                            "qualified": "users",
                            "columns": [
                                {"name": "username", "dtype": "text"},
                                {"name": "hashed_password", "dtype": "text"},
                            ],
                        }
                    ],
                }
            },
        },
        headers=headers,
    )
    assert forged.status_code == 200, forged.text
    assert "staging" not in forged.json()["connection_info"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "伪造 staging 库", "description": "", "data_source_id": forged.json()["id"]},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]

    for col in ("username", "hashed_password"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={
                "table_name": "users",
                "field_name": col,
                "alias_zh": col,
                "visible": True,
            },
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看用户密码哈希", "theme_ids": [theme_id]},
        headers=headers,
    )
    # Forged staging is stripped on create, so the theme is unqueryable and must
    # not fall back to demo rows (which would look like a successful analysis).
    assert chat.status_code == 422, chat.text
    assert chat.json()["detail"] == "Unable to query the selected themes."


def test_chat_ignores_persisted_forged_staging_metadata(client: TestClient):
    with SessionLocal() as db:
        user = db.query(User).filter(User.username == "admin").first()
        assert user

        ds = DataSource(
            name="persisted-forged-staging",
            source_type="excel",
            connection_info={},
            owner_id=user.id,
            status="active",
        )
        db.add(ds)
        db.commit()
        db.refresh(ds)

        ds.connection_info = {
            "staging": {
                "dialect": engine.dialect.name,
                "schema": None,
                "tables": [
                    {
                        "table": "users",
                        "schema": None,
                        "qualified": '"users"',
                        "columns": [
                            {"name": "username", "dtype": "text"},
                            {"name": "hashed_password", "dtype": "text"},
                        ],
                    }
                ],
            }
        }
        theme = ThemeLibrary(
            name="持久伪造 staging 库",
            description="",
            owner_id=user.id,
            data_source_id=ds.id,
        )
        db.add_all([ds, theme])
        db.commit()
        db.refresh(theme)
        db.add_all(
            [
                ThemeField(
                    theme_id=theme.id,
                    table_name="users",
                    field_name="username",
                    alias_zh="username",
                    visible=True,
                ),
                ThemeField(
                    theme_id=theme.id,
                    table_name="users",
                    field_name="hashed_password",
                    alias_zh="hashed_password",
                    visible=True,
                ),
            ]
        )
        db.commit()

        with pytest.raises(HTTPException) as exc:
            run_chat_query(db, user, "查看用户密码哈希", [theme.id])

    assert exc.value.status_code == 422
    assert exc.value.detail == "Unable to query the selected themes."


def test_sanitize_dataframe_columns_avoids_suffix_collisions():
    """
    Concrete upload crash: headers foo / foo_1 / foo! all need distinct
    physical names. The old per-base counter mapped the third column to
    foo_1 as well, so to_sql raised DuplicateColumnError.
    """
    df = pd.DataFrame({"foo": [1], "foo_1": [2], "foo!": [3]})
    out = _sanitize_dataframe_columns(df)
    assert list(out.columns) == ["foo", "foo_1", "foo_2"]
    assert not out.columns.duplicated().any()
    assert out["foo"].tolist() == [1]
    assert out["foo_1"].tolist() == [2]
    assert out["foo_2"].tolist() == [3]

    # Trailing-space / punctuation variants that Excel exports often produce.
    df2 = pd.DataFrame({"Revenue": [10], "Revenue_1": [20], "Revenue ": [30]})
    out2 = _sanitize_dataframe_columns(df2)
    assert list(out2.columns) == ["revenue", "revenue_1", "revenue_2"]
    assert not out2.columns.duplicated().any()


def test_excel_upload_survives_sanitized_column_name_collision(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    buf = io.BytesIO()
    pd.DataFrame({"foo": [1], "foo_1": [2], "foo!": [3]}).to_excel(
        buf, sheet_name="Sheet1", index=False, engine="openpyxl"
    )
    files = {
        "file": (
            "cols.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post("/api/data-sources/excel/upload", files=files, headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "active"
    names = [c["name"] for c in body["connection_info"]["staging"]["tables"][0]["columns"]]
    assert names == ["foo", "foo_1", "foo_2"]
    assert len(set(names)) == 3


def test_materialize_excel_keeps_colliding_column_values():
    buf = io.BytesIO()
    pd.DataFrame({"foo": [1], "foo_1": [2], "foo!": [3]}).to_excel(
        buf, sheet_name="Sheet1", index=False, engine="openpyxl"
    )
    info = materialize_excel_staging(engine, 99, buf.getvalue(), "cols.xlsx")
    table = info["staging"]["tables"][0]
    assert [c["name"] for c in table["columns"]] == ["foo", "foo_1", "foo_2"]
    from sqlalchemy import text

    qualified = table["qualified"]
    with engine.connect() as conn:
        row = conn.execute(text(f"SELECT foo, foo_1, foo_2 FROM {qualified}")).mappings().one()
    assert dict(row) == {"foo": 1, "foo_1": 2, "foo_2": 3}


def test_sanitize_dataframe_columns_fits_postgres_identifier_limit():
    """
    Residual of the suffix-collision fix: occupied-set uniquing produced
    `aaa…aaa` (63) and `aaa…aaa_1` (65). PostgreSQL NAMEDATALEN-1 truncates
    the second name back to 63, so CREATE TABLE raises DuplicateColumn.
    Live PG16 + psycopg: two headers of 63 and 64 A's → upload 500 and
    DataSource stuck status=staging (ValueError handler does not catch it).
    """
    long_a = "A" * MAX_SQL_IDENT_LEN
    long_b = "A" * (MAX_SQL_IDENT_LEN + 1)
    df = pd.DataFrame({long_a: [1], long_b: [2], "short": [3]})
    out = _sanitize_dataframe_columns(df)
    names = list(out.columns)
    assert len(names) == 3
    assert len(set(names)) == 3
    assert all(len(n) <= MAX_SQL_IDENT_LEN for n in names)
    # Simulated PG truncation must not reintroduce duplicates.
    truncated = [n[:MAX_SQL_IDENT_LEN] for n in names]
    assert len(set(truncated)) == 3
    assert out[names[0]].tolist() == [1]
    assert out[names[1]].tolist() == [2]
    assert out["short"].tolist() == [3]


def test_materialize_long_csv_filename_fits_postgres_identifier_limit():
    """
    CSV staging tables are ds_{id}_{sanitized_stem}. A descriptive export
    name makes that identifier exceed 63 characters; SQLAlchemy's PG dialect
    then raises IdentifierError (not ValueError) and upload 500s.
    """
    stem = "customer_transaction_history_export_north_america_q1_2024_final"
    filename = f"{stem}.csv"
    raw = b"district_name,amount\nA,1\nB,2\n"
    info = materialize_excel_staging(engine, 1, raw, filename)
    table = info["staging"]["tables"][0]
    assert len(table["table"]) <= MAX_SQL_IDENT_LEN
    assert table["table"].startswith("ds_1_")
    from sqlalchemy import text

    with engine.connect() as conn:
        rows = conn.execute(text(f"SELECT district_name, amount FROM {table['qualified']}")).mappings().all()
    assert [dict(r) for r in rows] == [
        {"district_name": "A", "amount": 1},
        {"district_name": "B", "amount": 2},
    ]


def test_materialize_preserves_literal_na_null_strings():
    """
    pandas default na_values treat 'NA' / 'NULL' / 'N/A' / 'None' as missing.
    Namibia's ISO code is NA; status columns often store those literals.
    They must land as text, not SQL NULL, or GROUP BY / filters drop the rows.
    """
    from sqlalchemy import text

    csv_raw = (
        b"country,status,revenue\n"
        b"NA,NULL,100\n"
        b"US,N/A,200\n"
        b"ZA,None,300\n"
        b"CN,#N/A,400\n"
    )
    info = materialize_excel_staging(engine, 401, csv_raw, "countries.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT country, status, revenue FROM {table['qualified']} ORDER BY revenue")
            ).mappings()
        ]
    assert rows == [
        {"country": "NA", "status": "NULL", "revenue": 100},
        {"country": "US", "status": "N/A", "revenue": 200},
        {"country": "ZA", "status": "None", "revenue": 300},
        {"country": "CN", "status": "#N/A", "revenue": 400},
    ]

    buf = io.BytesIO()
    pd.DataFrame(
        {
            "country": ["NA", "US", "ZA"],
            "status": ["NULL", "N/A", "None"],
            "revenue": [100, 200, 300],
        }
    ).to_excel(buf, sheet_name="Sheet1", index=False, engine="openpyxl")
    xinfo = materialize_excel_staging(engine, 402, buf.getvalue(), "countries.xlsx")
    xtable = xinfo["staging"]["tables"][0]
    with engine.connect() as conn:
        xrows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT country, status, revenue FROM {xtable['qualified']} ORDER BY revenue")
            ).mappings()
        ]
    assert xrows == [
        {"country": "NA", "status": "NULL", "revenue": 100},
        {"country": "US", "status": "N/A", "revenue": 200},
        {"country": "ZA", "status": "None", "revenue": 300},
    ]


def test_materialize_empty_numeric_cells_remain_null():
    """Empty amount cells must still become SQL NULL (not the string '') so SUM works."""
    from sqlalchemy import text

    raw = b"label,amount\nA,1.5\nB,\nC,3\n"
    info = materialize_excel_staging(engine, 403, raw, "amounts.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT label, amount FROM {table['qualified']} ORDER BY label")
            ).mappings()
        ]
        total = conn.execute(text(f"SELECT SUM(amount) AS s FROM {table['qualified']}")).scalar()
    assert rows == [
        {"label": "A", "amount": 1.5},
        {"label": "B", "amount": None},
        {"label": "C", "amount": 3.0},
    ]
    assert total == 4.5


def test_chat_query_preserves_literal_na_country_codes(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    raw = b"country,revenue\nNA,100\nUS,200\nNULL,50\n"
    up = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("na-countries.csv", raw, "text/csv")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "na-country-theme", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]
    for col in ("country", "revenue"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={"table_name": physical, "field_name": col, "alias_zh": col, "visible": True},
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看明细", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    countries = {row["country"] for row in chat.json()["rows"]}
    assert countries == {"NA", "US", "NULL"}
    assert None not in countries


def test_excel_upload_survives_postgres_length_column_collision(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    buf = io.BytesIO()
    pd.DataFrame({"A" * 63: [1], "A" * 64: [2]}).to_excel(
        buf, sheet_name="Sheet1", index=False, engine="openpyxl"
    )
    files = {
        "file": (
            "longcols.xlsx",
            buf.getvalue(),
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
    }
    r = client.post("/api/data-sources/excel/upload", files=files, headers=headers)
    assert r.status_code == 200, r.text
    body = r.json()
    assert body["status"] == "active"
    names = [c["name"] for c in body["connection_info"]["staging"]["tables"][0]["columns"]]
    assert len(names) == 2
    assert len(set(names)) == 2
    assert all(len(n) <= MAX_SQL_IDENT_LEN for n in names)


def test_excel_upload_marks_failed_when_materialize_raises_non_value_error(
    client: TestClient, monkeypatch
):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    def boom(*_args, **_kwargs):
        raise RuntimeError("duplicate column")

    monkeypatch.setattr(data_sources_routes, "materialize_excel_staging", boom)
    r = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("t.csv", b"a,b\n1,2\n", "text/csv")},
        headers=headers,
    )
    assert r.status_code == 500, r.text
    listed = client.get("/api/data-sources", headers=headers)
    assert listed.status_code == 200
    uploaded = [row for row in listed.json() if row["name"] == "t.csv"]
    assert uploaded, listed.text
    assert uploaded[0]["status"] == "failed"
    assert uploaded[0]["connection_info"] == {"error": "Failed to materialize upload."}


def test_materialize_preserves_leading_zero_identifiers():
    """
    pandas default inference turns 02101 / 000123 into int64 2101 / 123.
    Zip codes, bank accounts, and padded employee ids must stay text.
    """
    from sqlalchemy import text

    raw = b"zip,account_id,balance\n02101,000123,100.5\n02102,000124,200.0\n"
    info = materialize_excel_staging(engine, 501, raw, "accounts.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT zip, account_id, balance FROM {table['qualified']} ORDER BY zip")
            ).mappings()
        ]
        total = conn.execute(text(f"SELECT SUM(balance) AS s FROM {table['qualified']}")).scalar()
    assert rows == [
        {"zip": "02101", "account_id": "000123", "balance": 100.5},
        {"zip": "02102", "account_id": "000124", "balance": 200.0},
    ]
    assert total == 300.5

    wb = Workbook()
    ws = wb.active
    ws.append(["zip", "account_id", "balance"])
    for zip_code, account_id, balance in (("02101", "000123", 100.5), ("02102", "000124", 200.0)):
        ws.append([zip_code, account_id, balance])
        for col in ("A", "B"):
            ws[f"{col}{ws.max_row}"].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    xinfo = materialize_excel_staging(engine, 502, buf.getvalue(), "accounts.xlsx")
    xtable = xinfo["staging"]["tables"][0]
    with engine.connect() as conn:
        xrows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT zip, account_id, balance FROM {xtable['qualified']} ORDER BY zip")
            ).mappings()
        ]
    assert xrows == [
        {"zip": "02101", "account_id": "000123", "balance": 100.5},
        {"zip": "02102", "account_id": "000124", "balance": 200.0},
    ]


def test_materialize_large_ids_with_empty_cells_stay_text():
    """
    Integers in (2^53, 2^63) fit signed int64, so the previous coerce path
    treated them as numbers. Any empty cell makes pd.to_numeric return
    float64, which cannot represent those ids: 12345678901234567 and
    12345678901234568 both become 1.2345678901234568e16 and collide.
    Snowflake/order ids with optional blanks are a normal Excel export.
    """
    from sqlalchemy import text

    raw = (
        b"order_id,amount\n"
        b"12345678901234567,10.5\n"
        b",20\n"
        b"12345678901234568,30\n"
        b"9007199254740993,40\n"
    )
    info = materialize_excel_staging(engine, 504, raw, "snowflake.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(
                    f"SELECT order_id, amount FROM {table['qualified']} "
                    "ORDER BY amount"
                )
            ).mappings()
        ]
        total = conn.execute(text(f"SELECT SUM(amount) AS s FROM {table['qualified']}")).scalar()
    assert rows == [
        {"order_id": "12345678901234567", "amount": 10.5},
        {"order_id": None, "amount": 20.0},
        {"order_id": "12345678901234568", "amount": 30.0},
        {"order_id": "9007199254740993", "amount": 40.0},
    ]
    assert {row["order_id"] for row in rows if row["order_id"] is not None} == {
        "12345678901234567",
        "12345678901234568",
        "9007199254740993",
    }
    assert total == 100.5

    wb = Workbook()
    ws = wb.active
    ws.append(["order_id", "amount"])
    for order_id, amount in (
        ("12345678901234567", 10.5),
        (None, 20),
        ("12345678901234568", 30),
    ):
        ws.append([order_id, amount])
        if order_id is not None:
            ws[f"A{ws.max_row}"].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    xinfo = materialize_excel_staging(engine, 505, buf.getvalue(), "snowflake.xlsx")
    xtable = xinfo["staging"]["tables"][0]
    with engine.connect() as conn:
        xrows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT order_id, amount FROM {xtable['qualified']} ORDER BY amount")
            ).mappings()
        ]
    assert [row["order_id"] for row in xrows] == [
        "12345678901234567",
        None,
        "12345678901234568",
    ]


def test_materialize_large_decimal_ids_with_empty_cells_stay_text():
    """
    Decimal-typed id exports write 12345678901234567.0. The integer-only
    53-bit check treated any token with a '.' as a safe number, so empty
    cells still float64-collided those ids.
    """
    from sqlalchemy import text

    raw = (
        b"order_id,amount\n"
        b"12345678901234567.0,10.5\n"
        b",20\n"
        b"12345678901234568.00,30\n"
        b"-9007199254740993.0,40\n"
    )
    info = materialize_excel_staging(engine, 506, raw, "snowflake-decimal.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(
                    f"SELECT order_id, amount FROM {table['qualified']} "
                    "ORDER BY amount"
                )
            ).mappings()
        ]
        total = conn.execute(text(f"SELECT SUM(amount) AS s FROM {table['qualified']}")).scalar()
    assert rows == [
        {"order_id": "12345678901234567.0", "amount": 10.5},
        {"order_id": None, "amount": 20.0},
        {"order_id": "12345678901234568.00", "amount": 30.0},
        {"order_id": "-9007199254740993.0", "amount": 40.0},
    ]
    assert {row["order_id"] for row in rows if row["order_id"] is not None} == {
        "12345678901234567.0",
        "12345678901234568.00",
        "-9007199254740993.0",
    }
    assert total == 100.5

    wb = Workbook()
    ws = wb.active
    ws.append(["order_id", "amount"])
    for order_id, amount in (
        ("12345678901234567.0", 10.5),
        (None, 20),
        ("12345678901234568.0", 30),
    ):
        ws.append([order_id, amount])
        if order_id is not None:
            ws[f"A{ws.max_row}"].number_format = "@"
    buf = io.BytesIO()
    wb.save(buf)
    xinfo = materialize_excel_staging(engine, 507, buf.getvalue(), "snowflake-decimal.xlsx")
    xtable = xinfo["staging"]["tables"][0]
    with engine.connect() as conn:
        xrows = [
            dict(r)
            for r in conn.execute(
                text(
                    f"SELECT order_id, amount FROM {xtable['qualified']} "
                    "ORDER BY amount"
                )
            ).mappings()
        ]
    assert [row["order_id"] for row in xrows] == [
        "12345678901234567.0",
        None,
        "12345678901234568.0",
    ]


def test_materialize_large_uint64_ids_do_not_crash_upload():
    """
    Integers in (2^63, 2^64) become pandas uint64. SQLAlchemy to_sql then
    raises ValueError: Unsigned 64 bit integer datatype is not supported,
    so Excel/CSV upload 400s and the file never lands.
    """
    from sqlalchemy import text

    raw = b"order_id,amount\n12345678901234567890,10.5\n12345678901234567891,20\n"
    info = materialize_excel_staging(engine, 503, raw, "orders.csv")
    table = info["staging"]["tables"][0]
    with engine.connect() as conn:
        rows = [
            dict(r)
            for r in conn.execute(
                text(f"SELECT order_id, amount FROM {table['qualified']} ORDER BY order_id")
            ).mappings()
        ]
        total = conn.execute(text(f"SELECT SUM(amount) AS s FROM {table['qualified']}")).scalar()
    assert rows == [
        {"order_id": "12345678901234567890", "amount": 10.5},
        {"order_id": "12345678901234567891", "amount": 20.0},
    ]
    assert total == 30.5


def test_chat_query_preserves_large_ids_with_empty_cells(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    raw = b"order_id,revenue\n12345678901234567,100\n,50\n12345678901234568,200\n"
    up = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("snowflake-ids.csv", raw, "text/csv")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "snowflake-id-theme", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]
    for col in ("order_id", "revenue"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={"table_name": physical, "field_name": col, "alias_zh": col, "visible": True},
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看明细", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    ids = {row["order_id"] for row in chat.json()["rows"]}
    assert "12345678901234567" in ids
    assert "12345678901234568" in ids
    assert 12345678901234567 not in ids
    assert 1.2345678901234568e16 not in ids


def test_chat_query_preserves_large_decimal_ids_with_empty_cells(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    raw = b"order_id,revenue\n12345678901234567.0,100\n,50\n12345678901234568.0,200\n"
    up = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("snowflake-decimal-ids.csv", raw, "text/csv")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "snowflake-decimal-id-theme", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]
    for col in ("order_id", "revenue"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={"table_name": physical, "field_name": col, "alias_zh": col, "visible": True},
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看明细", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    ids = {row["order_id"] for row in chat.json()["rows"]}
    assert "12345678901234567.0" in ids
    assert "12345678901234568.0" in ids
    assert 1.2345678901234568e16 not in ids


def test_chat_query_preserves_leading_zero_zip_codes(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    raw = b"zip,revenue\n02101,100\n02102,200\n"
    up = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("zips.csv", raw, "text/csv")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "zip-leading-zero-theme", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]
    for col in ("zip", "revenue"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={"table_name": physical, "field_name": col, "alias_zh": col, "visible": True},
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看明细", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    zips = {row["zip"] for row in chat.json()["rows"]}
    assert zips == {"02101", "02102"}
    assert 2101 not in zips and "2101" not in zips


def test_chat_query_allows_update_column_name(client: TestClient):
    """Generated SQL quotes headers; the guard must not treat Update as UPDATE."""
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    headers = {"Authorization": f"Bearer {token}"}

    raw = b"Update,amount\nA,10\nB,20\n"
    up = client.post(
        "/api/data-sources/excel/upload",
        files={"file": ("updates.csv", raw, "text/csv")},
        headers=headers,
    )
    assert up.status_code == 200, up.text
    ds_id = up.json()["id"]
    physical = up.json()["connection_info"]["staging"]["tables"][0]["table"]

    theme = client.post(
        "/api/theme-libraries",
        json={"name": "update-column-theme", "description": "", "data_source_id": ds_id},
        headers=headers,
    )
    assert theme.status_code == 200, theme.text
    theme_id = theme.json()["id"]
    for col in ("update", "amount"):
        fr = client.post(
            f"/api/theme-libraries/{theme_id}/fields",
            json={"table_name": physical, "field_name": col, "alias_zh": col, "visible": True},
            headers=headers,
        )
        assert fr.status_code == 200, fr.text

    chat = client.post(
        "/api/chat/query",
        json={"prompt": "查看明细", "theme_ids": [theme_id]},
        headers=headers,
    )
    assert chat.status_code == 200, chat.text
    labels = {row["update"] for row in chat.json()["rows"]}
    assert labels == {"A", "B"}


def test_excel_upload_rejects_bad_extension(client: TestClient):
    login = client.post("/api/auth/login", json={"username": "admin", "password": "admin123"})
    token = login.json()["access_token"]
    files = {"file": ("bad.txt", b"hello", "text/plain")}
    r = client.post("/api/data-sources/excel/upload", files=files, headers={"Authorization": f"Bearer {token}"})
    assert r.status_code == 400


def test_excel_upload_reads_only_limit_plus_one(monkeypatch):
    class RecordingFile:
        read_size = None

        def read(self, size=-1):
            self.read_size = size
            return b"12345"

    class FakeUpload:
        filename = "large.csv"
        file = RecordingFile()

    monkeypatch.setattr(data_sources_routes, "MAX_UPLOAD_BYTES", 4)

    with pytest.raises(HTTPException) as exc:
        data_sources_routes.upload_excel(file=FakeUpload(), db=None, current_user=None)

    assert exc.value.status_code == 413
    assert FakeUpload.file.read_size == 5


def test_to_sql_chunksize_stays_under_postgres_bind_limit():
    """
    Production Docker uses postgresql+psycopg. method=multi with chunksize=500
    sends (rows * cols) binds per INSERT. psycopg then raises
    OperationalError: number of parameters must be between 0 and 65535,
    and the upload transaction rolls back.
    """
    assert _to_sql_chunksize(50, dialect="postgresql") == DEFAULT_TO_SQL_CHUNKSIZE
    assert _to_sql_chunksize(50, dialect="sqlite") == DEFAULT_TO_SQL_CHUNKSIZE

    wide = 200
    chunk = _to_sql_chunksize(wide, dialect="postgresql")
    assert chunk == 327
    assert chunk * wide <= PG_MAX_BIND_PARAMS
    assert (chunk + 1) * wide > PG_MAX_BIND_PARAMS

    # 132 columns * 500 rows = 66000 binds — the smallest typical failure.
    assert _to_sql_chunksize(132, dialect="postgresql") * 132 <= PG_MAX_BIND_PARAMS

    max_cols = 1600
    max_chunk = _to_sql_chunksize(max_cols, dialect="postgresql")
    assert max_chunk == 40
    assert max_chunk * max_cols <= PG_MAX_BIND_PARAMS


def _pg_test_engine():
    from sqlalchemy import create_engine, text

    url = os.environ.get(
        "TALKBI_PG_TEST_URL",
        "postgresql+psycopg://talkbi:talkbi@127.0.0.1:5432/talkbi",
    )
    eng = create_engine(url)
    try:
        with eng.connect() as conn:
            conn.execute(text("SELECT 1"))
    except Exception as exc:
        pytest.skip(f"PostgreSQL is not available: {exc}")
    return eng


def test_materialize_wide_csv_survives_postgres_bind_limit():
    """
    Concrete (psycopg on PostgreSQL 16, before fix):
    200-column × 400-row CSV → OperationalError bind overflow, no staging table.
    """
    from sqlalchemy import text

    header = ",".join(f"c{i}" for i in range(200))
    body = "\n".join(",".join("1" for _ in range(200)) for _ in range(400))
    raw = f"{header}\n{body}".encode()
    pg = _pg_test_engine()
    info = materialize_excel_staging(pg, 9101, raw, "wide.csv")
    table = info["staging"]["tables"][0]
    assert table["row_count"] == 400
    assert len(table["columns"]) == 200
    with pg.connect() as conn:
        n = conn.execute(text(f"SELECT COUNT(*) FROM {table['qualified']}")).scalar()
    assert n == 400
